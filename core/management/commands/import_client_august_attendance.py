from __future__ import annotations

import re
from datetime import datetime, date, time
from pathlib import Path
from zoneinfo import ZoneInfo

from django.core.management.base import BaseCommand, CommandError
from django.db import IntegrityError, transaction

from core.models import AttendanceRecord, Branch, FinalizedDTR, UserProfile


class Command(BaseCommand):
    help = (
        "Import the client's August-1-31-2026.xlsx workbook directly into "
        "AttendanceRecord. Reads the USER sheet for employee names and the "
        "Clock sheet for raw biometric punches. Runs as DRY-RUN by default."
    )

    MANILA = ZoneInfo("Asia/Manila")

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            required=True,
            help='Path to the client .xlsx file, e.g. "client_data/August-1-31-2026.xlsx"',
        )
        parser.add_argument(
            "--branch",
            default="DA MIMAROPA Client Validation",
            help="Branch name to attach the imported attendance to.",
        )
        parser.add_argument(
            "--ids",
            default="",
            help='Optional comma-separated biometric IDs, e.g. "27,19,20". Empty = all.',
        )
        parser.add_argument(
            "--commit",
            action="store_true",
            help="Actually write AttendanceRecord rows. Without this flag, only preview.",
        )
        parser.add_argument(
            "--replace-imported",
            action="store_true",
            help=(
                "Before importing, delete only rows previously created by THIS command "
                "for the same branch/date range. Does not delete manual/device rows."
            ),
        )

    def handle(self, *args, **options):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise CommandError(
                "openpyxl is required. Run: pip install openpyxl"
            )

        file_path = Path(options["file"]).expanduser().resolve()
        if not file_path.exists():
            raise CommandError(f"File not found: {file_path}")

        branch_name = str(options["branch"]).strip()
        if not branch_name:
            raise CommandError("Branch name cannot be blank.")

        selected_ids = {
            x.strip()
            for x in str(options["ids"] or "").split(",")
            if x.strip()
        }

        commit = bool(options["commit"])
        replace_imported = bool(options["replace_imported"])

        # IMPORTANT:
        # The client's "Clock" worksheet reports a formatted max_row close to
        # Excel's full 1,048,576-row limit even though the real employee data
        # ends much earlier.  openpyxl read_only=True reparses the XML for
        # repeated random cell access and becomes extremely slow on this file.
        #
        # Load the small (~hundreds of KB) workbook normally instead. This keeps
        # cells in memory and makes the bounded Clock-sheet scan finish quickly.
        wb = load_workbook(
            file_path,
            data_only=True,
            read_only=False,
        )

        required_sheets = {"Clock", "USER"}
        missing = required_sheets - set(wb.sheetnames)
        if missing:
            raise CommandError(
                "This is not the expected client workbook. Missing sheet(s): "
                + ", ".join(sorted(missing))
            )

        ws_clock = wb["Clock"]
        ws_user = wb["USER"]

        start_date, end_date = self._parse_workbook_range(ws_clock)
        roster = self._read_roster(ws_user)

        rows, employee_stats = self._read_clock_rows(
            ws_clock=ws_clock,
            roster=roster,
            start_date=start_date,
            end_date=end_date,
            selected_ids=selected_ids,
        )

        if not rows:
            raise CommandError(
                "No attendance punches were found for the selected biometric IDs."
            )

        self.stdout.write("")
        self.stdout.write(
            self.style.MIGRATE_HEADING(
                "CLIENT EXCEL ATTENDANCE IMPORT"
            )
        )
        self.stdout.write("=" * 78)
        self.stdout.write(f"Source file : {file_path}")
        self.stdout.write(f"Date range  : {start_date} to {end_date}")
        self.stdout.write(f"Branch      : {branch_name}")
        self.stdout.write(
            f"Mode        : {'COMMIT' if commit else 'DRY-RUN (nothing written)'}"
        )
        self.stdout.write(f"Punches     : {len(rows)}")
        self.stdout.write(
            f"Employees   : {len(employee_stats)} with at least one punch"
        )
        self.stdout.write(
            f"Workbook Clock reported max row: {ws_clock.max_row:,} "
            "(bounded importer scan used)"
        )
        self.stdout.write("")

        self.stdout.write("Employee preview:")
        for employee_id, stats in list(employee_stats.items())[:20]:
            self.stdout.write(
                f"  ID {employee_id:<8} "
                f"{stats['name'][:32]:<32} "
                f"Punches={stats['punches']:<3} "
                f"DaysWithLogs={len(stats['days']):<2}"
            )

        if len(employee_stats) > 20:
            self.stdout.write(
                f"  ... plus {len(employee_stats) - 20} more employee(s)"
            )

        if not commit:
            self.stdout.write("")
            self.stdout.write(
                self.style.WARNING(
                    "DRY-RUN ONLY. Review the counts above. "
                    "To actually import, rerun with --commit."
                )
            )
            self.stdout.write(
                'Example: python manage.py import_client_august_attendance '
                f'--file "{options["file"]}" --branch "{branch_name}" --commit'
            )
            return

        with transaction.atomic():
            branch, _ = Branch.objects.get_or_create(
                name=branch_name
            )

            self._check_locked_dtrs(
                branch=branch,
                start_date=start_date,
                end_date=end_date,
                employee_ids=set(employee_stats.keys()),
            )

            if replace_imported:
                deleted = self._delete_previous_import(
                    branch=branch,
                    start_date=start_date,
                    end_date=end_date,
                )
                self.stdout.write(
                    f"Previous rows from this importer removed: {deleted}"
                )

            created = 0
            skipped = 0

            for item in rows:
                try:
                    AttendanceRecord.objects.create(
                        employee_id=item["employee_id"],
                        full_name=item["full_name"],
                        department=item["department"],
                        branch=branch,
                        timestamp=item["timestamp"],
                        attendance_status=item["attendance_status"],
                        raw_row=item["raw_row"],
                    )
                    created += 1
                except IntegrityError:
                    skipped += 1

        self.stdout.write("")
        self.stdout.write("=" * 78)
        self.stdout.write(
            self.style.SUCCESS("CLIENT ATTENDANCE IMPORT COMPLETE")
        )
        self.stdout.write(
            self.style.SUCCESS(f"Created: {created}")
        )
        self.stdout.write(
            self.style.SUCCESS(f"Duplicates skipped: {skipped}")
        )
        self.stdout.write("")
        self.stdout.write(
            "Next: create or update the payroll profile for the employee(s) "
            "you want to test, using the SAME biometric ID. "
            "The system will then use these imported punches automatically."
        )

    def _parse_workbook_range(self, ws_clock):
        raw = ws_clock["F2"].value
        text = str(raw or "").strip()

        match = re.search(
            r"(\d{2})-(\d{2})-(\d{4})\s*~\s*"
            r"(\d{2})-(\d{2})-(\d{4})",
            text,
        )

        if not match:
            raise CommandError(
                f'Could not parse Clock!F2 date range: "{text}"'
            )

        sm, sd, sy, em, ed, ey = map(int, match.groups())

        return (
            date(sy, sm, sd),
            date(ey, em, ed),
        )

    def _read_roster(self, ws_user):
        """
        USER layout confirmed from client workbook:
          D = biometric/employee number
          E = biometric short name
          G = full employee name
          H = location/department label
          I = office/department
          K = designation
        """
        roster = {}

        for row in range(4, ws_user.max_row + 1):
            raw_id = ws_user.cell(row=row, column=4).value
            if raw_id in (None, ""):
                continue

            employee_id = self._clean_id(raw_id)
            if not employee_id:
                continue

            full_name = str(
                ws_user.cell(row=row, column=7).value or ""
            ).strip()

            bio_name = str(
                ws_user.cell(row=row, column=5).value or ""
            ).strip()

            office = str(
                ws_user.cell(row=row, column=9).value or ""
            ).strip()

            location = str(
                ws_user.cell(row=row, column=8).value or ""
            ).strip()

            designation = str(
                ws_user.cell(row=row, column=11).value or ""
            ).strip()

            roster[employee_id] = {
                "full_name": full_name or bio_name or f"Employee {employee_id}",
                "bio_name": bio_name,
                "department": office or location,
                "position": designation,
            }

        return roster

    def _read_clock_rows(
        self,
        ws_clock,
        roster,
        start_date,
        end_date,
        selected_ids,
    ):
        """
        Clock layout confirmed from client workbook:
          Employee header every 4 rows:
            D = "ID"
            F = biometric ID
            M = "Name"
            O = biometric short name

          header+1: day numbers in D:AH
          header+2: weekday labels
          header+3: newline-separated punches in D:AH
        """
        rows = []
        employee_stats = {}

        current_row = 1

        # Do NOT scan ws_clock.max_row.  In the client workbook it is almost
        # 1,048,576 because of formatting, while the real Clock data is only a
        # few hundred rows.  The USER roster gives us a safe upper bound:
        # each Clock employee occupies 4 rows, plus a small header/buffer.
        scan_limit = min(
            ws_clock.max_row,
            max(100, (len(roster) * 4) + 50),
        )

        while current_row <= scan_limit:
            marker = ws_clock.cell(
                row=current_row,
                column=4,
            ).value

            if str(marker or "").strip().upper() != "ID":
                current_row += 1
                continue

            raw_id = ws_clock.cell(
                row=current_row,
                column=6,
            ).value

            employee_id = self._clean_id(raw_id)

            if not employee_id:
                current_row += 4
                continue

            if selected_ids and employee_id not in selected_ids:
                current_row += 4
                continue

            short_name = str(
                ws_clock.cell(
                    row=current_row,
                    column=15,
                ).value
                or ""
            ).strip()

            person = roster.get(
                employee_id,
                {
                    "full_name": short_name or f"Employee {employee_id}",
                    "bio_name": short_name,
                    "department": "",
                    "position": "",
                },
            )

            day_row = current_row + 1
            log_row = current_row + 3

            stats = {
                "name": person["full_name"],
                "punches": 0,
                "days": set(),
            }

            # D through AH = days 1 through 31.
            for col in range(4, 35):
                day_value = ws_clock.cell(
                    row=day_row,
                    column=col,
                ).value

                try:
                    day_num = int(day_value)
                except (TypeError, ValueError):
                    continue

                try:
                    current_date = date(
                        start_date.year,
                        start_date.month,
                        day_num,
                    )
                except ValueError:
                    continue

                if not (
                    start_date
                    <= current_date
                    <= end_date
                ):
                    continue

                raw_cell = ws_clock.cell(
                    row=log_row,
                    column=col,
                ).value

                if raw_cell in (None, ""):
                    continue

                punches = self._parse_punches(raw_cell)

                if not punches:
                    continue

                stats["days"].add(current_date)

                for index, punch_time in enumerate(punches):
                    attendance_status = (
                        AttendanceRecord.STATUS_CHECKIN
                        if index % 2 == 0
                        else AttendanceRecord.STATUS_CHECKOUT
                    )

                    timestamp = datetime(
                        current_date.year,
                        current_date.month,
                        current_date.day,
                        punch_time.hour,
                        punch_time.minute,
                        punch_time.second,
                        tzinfo=self.MANILA,
                    )

                    rows.append(
                        {
                            "employee_id": employee_id,
                            "full_name": person["full_name"],
                            "department": person["department"],
                            "timestamp": timestamp,
                            "attendance_status": attendance_status,
                            "raw_row": {
                                "source": "August-1-31-2026.xlsx",
                                "importer": "import_client_august_attendance_v1",
                                "sheet": "Clock",
                                "clock_cell": ws_clock.cell(
                                    row=log_row,
                                    column=col,
                                ).coordinate,
                                "original_clock_value": str(raw_cell),
                                "punch_sequence": index + 1,
                                "biometric_short_name": short_name,
                                "position_from_user_sheet": person["position"],
                            },
                        }
                    )

                    stats["punches"] += 1

            if stats["punches"] > 0:
                employee_stats[employee_id] = stats

            current_row += 4

        return rows, employee_stats

    def _parse_punches(self, raw_value):
        text = str(raw_value or "").strip()
        if not text:
            return []

        # Client workbook stores times separated by line breaks.
        # Also accept spaces/semicolons as a fallback.
        tokens = re.findall(
            r"\b([01]?\d|2[0-3]):([0-5]\d)(?::([0-5]\d))?\b",
            text,
        )

        punches = []

        for hh, mm, ss in tokens:
            punches.append(
                time(
                    int(hh),
                    int(mm),
                    int(ss or 0),
                )
            )

        # Preserve the source order; the Clock sheet is chronological.
        return punches

    def _clean_id(self, value):
        if value in (None, ""):
            return ""

        text = str(value).strip()

        # Excel sometimes presents integer IDs as "27.0".
        if re.fullmatch(r"\d+\.0+", text):
            text = text.split(".", 1)[0]

        return text.lstrip("'").strip()

    def _check_locked_dtrs(
        self,
        branch,
        start_date,
        end_date,
        employee_ids,
    ):
        profiles = UserProfile.objects.filter(
            branch=branch,
            biometric_employee_id__in=employee_ids,
        )

        conflicts = []

        for profile in profiles:
            locked = FinalizedDTR.objects.filter(
                profile=profile,
                branch=branch,
                is_locked=True,
                period__start_date__lte=end_date,
                period__end_date__gte=start_date,
            ).select_related("period").first()

            if locked:
                conflicts.append(
                    f"{profile.user.username}: {locked.period.name}"
                )

        if conflicts:
            raise CommandError(
                "Import blocked because these employees have a locked DTR "
                "overlapping the client workbook period: "
                + "; ".join(conflicts[:10])
            )

    def _delete_previous_import(
        self,
        branch,
        start_date,
        end_date,
    ):
        qs = AttendanceRecord.objects.filter(
            branch=branch,
            timestamp__date__gte=start_date,
            timestamp__date__lte=end_date,
            raw_row__importer="import_client_august_attendance_v1",
        )

        count = qs.count()
        qs.delete()
        return count
